from __future__ import annotations

import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

from dailyscraper.domain import DeliveryRecord, DeliveryStatus
from dailyscraper.execution.partitions import ItemIdentity, PartitionSink, PartitionWriter
from dailyscraper.output import (
    Deliverer,
    DeliveryContext,
    iter_csv,
    make_archive,
    scraper_rows,
    write_output,
)
from dailyscraper.schema import Inputs
from dailyscraper.settings.models import DeliverySettings
from dailyscraper.storage import RunLayout
from dailyscraper.storage.sqlite import SQLiteStore
from dailyscraper.tasks.inputs import create_input_snapshot

ROWS = [
    {"name": "Чайник; белый", "price": "10", "tags": ["a", "b"]},
    {"name": 'Quote "x" <y>', "extra": None},
]


def test_csv_has_all_columns_in_first_seen_order(tmp_path: Path) -> None:
    path = tmp_path / "out.csv"

    assert write_output("csv", lambda: ROWS, path) == 2

    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    rows = list(iter_csv(path))
    assert list(rows[0]) == ["name", "price", "tags", "extra"]
    assert rows[0]["name"] == "Чайник; белый"
    assert rows[0]["tags"] == '["a", "b"]'
    assert rows[1] == {"name": 'Quote "x" <y>', "price": "", "tags": "", "extra": ""}


def test_json_keeps_values(tmp_path: Path) -> None:
    path = tmp_path / "out.json"
    write_output("json", lambda: ROWS, path)

    assert json.loads(path.read_text(encoding="utf-8")) == ROWS

    empty = tmp_path / "empty.json"
    write_output("json", lambda: [], empty)
    assert json.loads(empty.read_text(encoding="utf-8")) == []


def test_xlsx_has_header_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    write_output("xlsx", lambda: ROWS, path)

    sheet = load_workbook(path).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0] == ["name", "price", "tags", "extra"]
    assert values[1][0] == "Чайник; белый"
    assert len(values) == 3


def test_xml_escapes_values(tmp_path: Path) -> None:
    path = tmp_path / "out.xml"
    write_output("xml", lambda: ROWS, path)

    rows = ElementTree.parse(path).getroot().findall("row")
    assert rows[1].find("field[@name='name']").text == 'Quote "x" <y>'


def test_unknown_format_and_failed_write_leave_no_file(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        write_output("yaml", lambda: ROWS, tmp_path / "out.yaml")

    def broken():
        yield {"a": 1}
        raise RuntimeError("source failed")

    with pytest.raises(RuntimeError):
        write_output("json", broken, tmp_path / "out.json")
    assert list(tmp_path.iterdir()) == []


def test_archive_contains_the_files(tmp_path: Path) -> None:
    first = tmp_path / "a.csv"
    first.write_text("a", encoding="utf-8")
    second = tmp_path / "b.json"
    second.write_text("[]", encoding="utf-8")

    archive = make_archive([first, second], tmp_path / "out" / "run.zip")

    assert sorted(zipfile.ZipFile(archive).namelist()) == ["a.csv", "b.json"]


def make_run(tmp_path: Path):
    store = SQLiteStore(tmp_path / "store.sqlite")
    layout = RunLayout(tmp_path / "run")
    sink = PartitionSink(
        PartitionWriter(
            layout.root,
            run_id="r1",
            scraper="details",
            writer_id="w1",
            on_commit=store.partitions.commit,
            on_seal=store.partitions.seal,
            fsync=False,
        ),
        identity=ItemIdentity(),
        batch_rows=1,
    )
    return store, layout, sink


def test_list_rows_follow_commit_order(tmp_path: Path) -> None:
    store, layout, sink = make_run(tmp_path)
    for name in ("b", "a", "c"):
        sink.emit({"name": name})
    sink.close()

    rows = list(scraper_rows(store.partitions, layout, "r1", "details"))

    assert [row["name"] for row in rows] == ["b", "a", "c"]


def test_detail_rows_follow_input_order_with_latest_result(tmp_path: Path) -> None:
    store, layout, sink = make_run(tmp_path)
    source = tmp_path / "input.csv"
    source.write_text("id;link\n1;https://a/1\n2;https://a/2\n3;https://a/3\n", encoding="utf-8")
    snapshot = create_input_snapshot(
        source, layout.inputs / "details.csv", Inputs(required_fields=("link",))
    )
    sink.emit({"id": "3", "title": "third"}, input_row=3)
    sink.emit({"id": "1", "title": "first, try 1"}, input_row=1, status="timeout")
    sink.emit({"id": "1", "title": "first, try 2"}, input_row=1)
    sink.close()

    rows = list(scraper_rows(store.partitions, layout, "r1", "details", snapshot))

    assert rows == [
        {"id": "1", "title": "first, try 2"},
        {"id": "2", "link": "https://a/2"},
        {"id": "3", "title": "third"},
    ]


def context(tmp_path: Path, *, succeeded: bool = True) -> DeliveryContext:
    return DeliveryContext(
        run_id="r1", project_id="default", date="2026-10-28", succeeded=succeeded, work_dir=tmp_path
    )


def outputs(tmp_path: Path) -> dict[str, list[Path]]:
    csv_file = tmp_path / "catalog.csv"
    csv_file.write_text("a;b\n", encoding="utf-8")
    json_file = tmp_path / "catalog.json"
    json_file.write_text("[]", encoding="utf-8")
    return {"csv": [csv_file], "json": [json_file]}


def deliverer(tmp_path: Path, records: list[DeliveryRecord], **kwargs) -> Deliverer:
    return Deliverer(
        base_dir=tmp_path,
        clock=lambda: "2026-10-28T10:00:00+00:00",
        record=records.append,
        sleep=lambda seconds: None,
        **kwargs,
    )


def test_folder_delivery_copies_selected_formats(tmp_path: Path) -> None:
    records: list[DeliveryRecord] = []
    settings = DeliverySettings(
        name="exports", type="folder", path="exports/{date}/{run_id}", formats=("csv",)
    )

    (result,) = deliverer(tmp_path, records).run([settings], outputs(tmp_path), context(tmp_path))

    target = tmp_path / "exports" / "2026-10-28" / "r1"
    assert result.status is DeliveryStatus.DELIVERED
    assert [path.name for path in target.iterdir()] == ["catalog.csv"]
    assert records == [result]


def test_archived_delivery_sends_one_zip(tmp_path: Path) -> None:
    settings = DeliverySettings(name="zip", type="folder", path="out", archive=True)

    (result,) = deliverer(tmp_path, []).run([settings], outputs(tmp_path), context(tmp_path))

    assert result.files == ("r1-zip.zip",)
    archive = zipfile.ZipFile(tmp_path / "out" / "r1-zip.zip")
    assert sorted(archive.namelist()) == ["catalog.csv", "catalog.json"]


def test_success_only_delivery_is_skipped_after_a_failed_run(tmp_path: Path) -> None:
    settings = DeliverySettings(name="exports", type="folder", path="out")

    (result,) = deliverer(tmp_path, []).run(
        [settings], outputs(tmp_path), context(tmp_path, succeeded=False)
    )

    assert result.status is DeliveryStatus.SKIPPED
    assert not (tmp_path / "out").exists()


def test_failed_delivery_is_retried_and_recorded(tmp_path: Path) -> None:
    calls = []

    class Flaky:
        kind = "folder"

        def deliver(self, files, ctx):
            calls.append(files)
            if len(calls) < 3:
                raise OSError("network share offline")
            return "share"

    settings = DeliverySettings(name="share", type="folder", path="x", attempts=3)
    (ok,) = deliverer(tmp_path, [], destinations=lambda s, base: Flaky()).run(
        [settings], outputs(tmp_path), context(tmp_path)
    )
    assert (ok.status, ok.attempts, ok.location) == (DeliveryStatus.DELIVERED, 3, "share")

    calls.clear()
    settings = DeliverySettings(name="share", type="folder", path="x", attempts=2)
    (failed,) = deliverer(tmp_path, [], destinations=lambda s, base: Flaky()).run(
        [settings], outputs(tmp_path), context(tmp_path)
    )
    assert failed.status is DeliveryStatus.FAILED
    assert "network share offline" in failed.error
